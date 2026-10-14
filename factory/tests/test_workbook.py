import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from helpers import EXAMPLE
from unit_economics import load_config
from unit_economics.verify import compare, recalculate, soffice_path, verify_workbook
from unit_economics.workbook import build_workbook

# Functions LibreOffice cannot evaluate, or that need special handling in files written by openpyxl.
UNSAFE = re.compile(r"\b(XLOOKUP|XMATCH|FILTER|SORT|UNIQUE|SEQUENCE|LET|LAMBDA|TEXTJOIN|IFS|SWITCH|MAXIFS|MINIFS)\(")


class StructureTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb, cls.cells, cls.checks, cls.statuses = build_workbook(load_config(EXAMPLE))

    def test_sheets(self):
        self.assertEqual(
            self.wb.sheetnames,
            ["README", "Inputs", "Offline", "Online Web store", "Ladder Web store", "Scenarios", "Checks"],
        )

    def test_named_inputs(self):
        for name in ("MRP", "LandedCost", "GSTRate", "AllInCost", "LoadedOffline", "LoadedOnline", "TgtOnlineContribution"):
            self.assertIn(name, self.wb.defined_names)

    def test_every_tracked_result_is_a_formula(self):
        self.assertGreater(len(self.cells), 150)
        for key, (sheet, coord) in self.cells.items():
            if key.startswith("scenarios.") and key.endswith(".units"):
                continue  # units is an input, tracked so the engine value can be compared
            value = self.wb[sheet][coord].value
            self.assertTrue(isinstance(value, str) and value.startswith("="), f"{key} at {sheet}!{coord} is {value!r}")

    def test_no_conditional_number_formats(self):
        # LibreOffice drops or doubles the minus sign inside conditional sections, so a loss could show as a profit
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotRegex(cell.number_format, r"\[[<>=]", f"{ws.title}!{cell.coordinate}")

    def test_only_portable_functions(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        self.assertIsNone(UNSAFE.search(cell.value), f"{ws.title}!{cell.coordinate}: {cell.value}")


@unittest.skipUnless(soffice_path(), "LibreOffice is not installed; formula recalculation tests skipped")
class RecalculationTest(unittest.TestCase):
    def test_workbook_matches_engine(self):
        report = verify_workbook(load_config(EXAMPLE))
        self.assertEqual(report.errors, [])
        self.assertEqual(report.mismatches, [])
        self.assertGreater(report.compared, 150)
        self.assertEqual(report.checks_compared, 10)

    def test_a_wrong_formula_is_caught(self):
        cfg = load_config(EXAMPLE)
        wb, cells, checks, _ = build_workbook(cfg)
        sheet, coord = cells["online.web.trio.contribution"]
        wb[sheet][coord] = "=1"
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "tampered.xlsx"
            wb.save(src)
            report = compare(cfg, recalculate(src, Path(tmp) / "out"), cells, checks)
        self.assertFalse(report.ok)
        self.assertTrue(any("online.web.trio.contribution" in m for m in report.mismatches))

    def test_recalculated_file_holds_values(self):
        cfg = load_config(EXAMPLE)
        wb, cells, _, _ = build_workbook(cfg)
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "m.xlsx"
            wb.save(src)
            values = load_workbook(recalculate(src, Path(tmp) / "out"), data_only=True)
            sheet, coord = cells["scenarios.month.operating_contribution"]
            self.assertEqual(values[sheet][coord].value, 16000)


if __name__ == "__main__":
    unittest.main()
