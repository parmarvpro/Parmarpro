"""Generate an .xlsx model with live formulas from a config.

The workbook mirrors the engine so the model can be used without Python:
inputs are blue (with a yellow fill while their status is still an
assumption) and every result is a formula. ``verify.py`` recalculates the
file in LibreOffice and checks each result cell against the engine.

Structural changes (adding an offer, route, cost line or scenario) are made in
the TOML config and the workbook is regenerated; values can be edited freely.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from .config import BASES, KINDS, STATUSES, UNIT_COST_BASES, VIA, Config, OnlineChannel, Param, Scenario

FONT = "Arial"
BLUE, BLACK, GREEN, WHITE, GREY, RED = "0000FF", "000000", "008000", "FFFFFF", "595959", "C00000"
SECTION_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
HEADER_BORDER = Border(bottom=Side(style="thin", color="1F3864"))
ROW_RULE = Border(bottom=Side(style="hair", color="BFBFBF"))

INR = '"₹"#,##0.00;("₹"#,##0.00);"-"'
# No conditional sections: spreadsheet apps disagree on the minus sign inside them, and a loss
# must never display as a profit. Excel applies lakh grouping to #,##0 on an Indian locale.
INR_BIG = '"₹"#,##0;("₹"#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM = "0.00"
INT = "#,##0"
UNVERIFIED = ("report", "assumption", "unspecified")
EXTRA_ROWS = 2  # blank input rows added to additive tables (components, opex, mix)


def _font(color: str = BLACK, bold: bool = False, size: int = 10, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def _abs(col: int, row: int) -> str:
    return f"${get_column_letter(col)}${row}"


def _rel(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _ext(sheet: str, col: int, row: int) -> str:
    return f"'{sheet}'!{_abs(col, row)}"


def _sheet_title(text: str, taken: set[str]) -> str:
    clean = "".join(ch for ch in text if ch not in "[]:*?/\\'").strip()[:31] or "Sheet"
    title, n = clean, 2
    while title in taken:
        suffix = f" {n}"
        title, n = clean[: 31 - len(suffix)] + suffix, n + 1
    taken.add(title)
    return title


class _Builder:
    def __init__(self, config: Config) -> None:
        self.cfg = config
        self.sym = config.meta.currency_symbol
        self.inr = INR.replace("₹", self.sym)
        self.inr_big = INR_BIG.replace("₹", self.sym)
        self.wb = Workbook()
        self.cells: dict[str, tuple[str, str]] = {}
        self.status_ranges: list[tuple[str, str]] = []
        self.check_cells: dict[str, tuple[str, str]] = {}
        self.taken: set[str] = set()
        self.dv_cache: dict[tuple[str, str], DataValidation] = {}

    # ----------------------------------------------------------------- primitives

    def sheet(self, title: str):
        title = _sheet_title(title, self.taken)
        if self.wb.active is not None and self.wb.active.title == "Sheet" and not self.wb.active["A1"].value:
            ws = self.wb.active
            ws.title = title
        else:
            ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        return ws

    def title(self, ws, text: str, sub: str = "") -> None:
        ws["A1"] = text
        ws["A1"].font = _font(bold=True, size=14)
        if sub:
            ws["A2"] = sub
            ws["A2"].font = _font(GREY, italic=True)

    def section(self, ws, row: int, text: str, width: int) -> None:
        for col in range(1, width + 1):
            ws.cell(row, col).fill = SECTION_FILL
        ws.cell(row, 1, text).font = _font(WHITE, bold=True, size=11)

    def header(self, ws, row: int, labels: list[str], start: int = 1) -> None:
        for i, label in enumerate(labels):
            cell = ws.cell(row, start + i, label)
            cell.font = _font(bold=True)
            cell.fill = HEADER_FILL
            cell.border = HEADER_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    def label(self, ws, row: int, col: int, text: str, bold: bool = False, color: str = BLACK) -> None:
        cell = ws.cell(row, col, text)
        cell.font = _font(color, bold=bold)

    def value(self, ws, row: int, col: int, value, fmt: str | None = None):
        cell = ws.cell(row, col, value)
        cell.font = _font(BLUE)
        if fmt:
            cell.number_format = fmt
        return cell

    def formula(self, ws, row: int, col: int, text: str, fmt: str | None = None, link: bool = False, bold: bool = False):
        cell = ws.cell(row, col, text)
        cell.font = _font(GREEN if link else BLACK, bold=bold)
        if fmt:
            cell.number_format = fmt
        return cell

    def validation(self, ws, choices: tuple[str, ...] | list[str], coord: str) -> None:
        key = (ws.title, ",".join(choices))
        dv = self.dv_cache.get(key)
        if dv is None:
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            self.dv_cache[key] = dv
        dv.add(coord)

    def status_cells(self, ws, row: int, status_col: int, param: Param | None) -> None:
        if param is None:
            return
        cell = ws.cell(row, status_col, param.status)
        cell.font = _font(BLUE)
        self.validation(ws, STATUSES, cell.coordinate)
        src = ws.cell(row, status_col + 1, param.source)
        src.font = _font(BLUE)

    def flag_unverified(self, ws, value_range: str, status_ref: str) -> None:
        """Yellow fill on `value_range` while the status cell says it is still an assumption."""
        test = ",".join(f'{status_ref}="{s}"' for s in UNVERIFIED)
        ws.conditional_formatting.add(value_range, FormulaRule(formula=[f"OR({test})"], fill=ASSUMPTION_FILL))

    def rule_rows(self, ws, first: int, last: int, ncols: int) -> None:
        """Hairline under each input row so blank rows read as places to type."""
        for row in range(first, last + 1):
            for col in range(1, ncols + 1):
                ws.cell(row, col).border = ROW_RULE

    def red_flags(self, ws, cell_range: str, top_left: str, words: tuple[str, ...]) -> None:
        test = ",".join(f'{top_left}="{w}"' for w in words)
        ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f"OR({test})"], font=Font(name=FONT, bold=True, color=RED)))

    def name(self, name: str, ws, row: int, col: int) -> None:
        self.wb.defined_names[name] = DefinedName(name, attr_text=_ext(ws.title, col, row))

    def track(self, key: str, ws, row: int, col: int) -> None:
        self.cells[key] = (ws.title, _rel(col, row))

    def widths(self, ws, widths: dict[str, float]) -> None:
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

    # ----------------------------------------------------------------- sheets

    def build(self) -> Workbook:
        readme = self.sheet("README")
        self.inputs_sheet()
        if self.cfg.routes:
            self.offline_sheet()
        for channel in self.cfg.online:
            self.online_sheet(channel)
        if self.cfg.scenarios:
            self.scenario_sheet()
        self.checks_sheet()
        self.readme_sheet(readme)
        for ws in self.wb.worksheets:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)
        return self.wb

    def inputs_sheet(self) -> None:
        cfg, inr = self.cfg, self.inr
        ws = self.sheet("Inputs")
        self.title(ws, "Inputs — product, cost build-up and targets", "Blue = input · yellow fill = still an assumption · black = formula")
        self.widths(ws, {"A": 46, "B": 14, "C": 13, "D": 13, "E": 13, "F": 70})
        unit = cfg.meta.unit
        p = cfg.product

        row = 4
        self.section(ws, row, "Product and tax", 6)
        self.header(ws, row + 1, ["Parameter", "Value", "Unit", "", "Status", "Source"])
        rows = [
            ("MRP (consumer price, tax-inclusive)", p.mrp, inr, f"per {unit}", "MRP"),
            ("Landed product cost", p.landed_cost, inr, f"per {unit}", "LandedCost"),
            ("Output GST rate on the product", p.gst_rate, PCT, "rate", "GSTRate"),
            ("Planning cost range — minimum", p.planning_cost_min, inr, f"per {unit}", "PlanMin"),
            ("Planning cost range — maximum", p.planning_cost_max, inr, f"per {unit}", "PlanMax"),
        ]
        row += 2
        first = row
        for label, param, fmt, unit_text, name in rows:
            self.label(ws, row, 1, label)
            self.value(ws, row, 2, param.value if param else None, fmt)
            self.label(ws, row, 3, unit_text, color=GREY)
            self.status_cells(ws, row, 5, param)
            self.name(name, ws, row, 2)
            row += 1
        self.status_ranges.append((ws.title, f"E{first}:E{row - 1}"))
        self.flag_unverified(ws, f"B{first}:B{row - 1}", f"$E{first}")
        self.rule_rows(ws, first, row - 1, 6)

        row += 1
        self.section(ws, row, f"Cost build-up — per {unit}, added to the landed cost", 6)
        self.header(ws, row + 1, ["Component", "Amount", "Offline (1/0)", "Online (1/0)", "Status", "Source"])
        row += 2
        comp_first = row
        for comp in p.components:
            self.label(ws, row, 1, comp.name)
            self.value(ws, row, 2, comp.amount.value, inr)
            for col, channel in ((3, "offline"), (4, "online")):
                self.value(ws, row, col, 1 if channel in comp.channels else 0, "0").alignment = Alignment(horizontal="center")
                self.validation(ws, ("0", "1"), _rel(col, row))
            self.status_cells(ws, row, 5, comp.amount)
            row += 1
        for _ in range(EXTRA_ROWS):
            for col, fmt in ((1, None), (2, inr), (3, "0"), (4, "0")):
                self.value(ws, row, col, None, fmt)
            for col in (3, 4):
                ws.cell(row, col).alignment = Alignment(horizontal="center")
                self.validation(ws, ("0", "1"), _rel(col, row))
            self.validation(ws, STATUSES, _rel(5, row))
            row += 1
        comp_last = row - 1
        self.rule_rows(ws, comp_first, comp_last, 6)
        ws.cell(row, 1, "Blank rows above take extra components; put 1 under each channel the cost applies to.").font = _font(GREY, italic=True)
        row += 2
        self.status_ranges.append((ws.title, f"E{comp_first}:E{comp_last}"))
        self.flag_unverified(ws, f"B{comp_first}:B{comp_last}", f"$E{comp_first}")
        amounts = f"$B${comp_first}:$B${comp_last}"

        summary = [
            ("Components total", f"=SUM({amounts})", inr, "ComponentsTotal", "cost.components_total"),
            ("All-in planning cost (landed + all components)", "=LandedCost+ComponentsTotal", inr, "AllInCost", "cost.all_in"),
            (
                "Loaded cost — offline routes",
                f"=LandedCost+SUMPRODUCT({amounts},$C${comp_first}:$C${comp_last})",
                inr,
                "LoadedOffline",
                "cost.loaded.offline",
            ),
            (
                "Loaded cost — online channels (if basis = loaded)",
                f"=LandedCost+SUMPRODUCT({amounts},$D${comp_first}:$D${comp_last})",
                inr,
                "LoadedOnline",
                "cost.loaded.online",
            ),
        ]
        for label, text, fmt, name, key in summary:
            self.label(ws, row, 1, label, bold=name == "AllInCost")
            self.formula(ws, row, 2, text, fmt, bold=name == "AllInCost")
            self.name(name, ws, row, 2)
            self.track(key, ws, row, 2)
            row += 1
        self.label(ws, row, 1, "All-in cost within the planning range?")
        self.formula(
            ws,
            row,
            2,
            '=IF(COUNT(PlanMin,PlanMax)<2,"n/a",IF(AND(AllInCost>=PlanMin-1E-9,AllInCost<=PlanMax+1E-9),"Yes","NO"))',
        )
        self.name("PlanCheck", ws, row, 2)
        row += 2

        t = cfg.targets
        self.section(ws, row, "Targets (leave blank to switch a check off)", 6)
        self.header(ws, row + 1, ["Target", "Value", "Unit", "", "Status", "Source"])
        row += 2
        targets = [
            ("Retailer margin — minimum (share of MRP)", t.retailer_margin_min, PCT, "of MRP", "TgtRetailerMarginMin"),
            ("Retailer margin — maximum (share of MRP)", t.retailer_margin_max, PCT, "of MRP", "TgtRetailerMarginMax"),
            ("Minimum gross margin on landed cost", t.min_gross_margin, PCT, "of net realization", "TgtGrossMargin"),
            ("Minimum offline contribution after loaded cost", t.min_offline_contribution, inr, f"per {unit}", "TgtOfflineContribution"),
            ("Minimum online contribution after acquisition", t.min_online_contribution, inr, "per order", "TgtOnlineContribution"),
        ]
        first = row
        for label, param, fmt, unit_text, name in targets:
            self.label(ws, row, 1, label)
            self.value(ws, row, 2, param.value if param else None, fmt)
            self.label(ws, row, 3, unit_text, color=GREY)
            if param:
                self.status_cells(ws, row, 5, param)
            else:
                self.validation(ws, STATUSES, _rel(5, row))
            self.name(name, ws, row, 2)
            row += 1
        self.status_ranges.append((ws.title, f"E{first}:E{row - 1}"))
        ws.freeze_panes = "A4"

    def offline_sheet(self) -> None:
        cfg, inr = self.cfg, self.inr
        unit = cfg.meta.unit
        routes = cfg.routes
        n = len(routes)
        ws = self.sheet("Offline")
        self.title(
            ws,
            "Offline routes — partner margins and brand contribution",
            f"Each column is one route. Brand realization is per {unit} shipped, after any free-goods scheme (buy X get Y free).",
        )
        self.widths(ws, {"A": 52, **{get_column_letter(2 + i): 18 for i in range(n)}})
        last_col = get_column_letter(1 + n)
        row = 4
        self.section(ws, row, "Route inputs", 1 + n)
        row += 1
        self.header(ws, row, ["Route"] + [r.name for r in routes])
        ws.row_dimensions[row].height = 30
        row += 1
        r: dict[str, int] = {}

        def put(key: str, label: str, bold: bool = False) -> None:
            nonlocal row
            self.label(ws, row, 1, label, bold=bold)
            r[key] = row
            row += 1

        put("via", "Via (direct / distributor)")
        put("price", f"Brand price per {unit} (to its direct customer)")
        put("retailer_price", "Retailer price (distributor routes only)")
        put("buy", "Free-goods scheme: buy")
        put("free", "Free-goods scheme: get free")
        put("status", "Price status")
        put("source", "Price source")
        for i, route in enumerate(routes):
            c = 2 + i
            self.value(ws, r["via"], c, route.via)
            self.validation(ws, VIA, _rel(c, r["via"]))
            self.value(ws, r["price"], c, route.price.value, inr)
            self.value(ws, r["retailer_price"], c, route.retailer_price.value if route.retailer_price else None, inr)
            self.value(ws, r["buy"], c, route.scheme_buy or None, "0")
            self.value(ws, r["free"], c, route.scheme_free or None, "0")
            status = ws.cell(r["status"], c, route.price.status)
            status.font = _font(BLUE)
            self.validation(ws, STATUSES, status.coordinate)
            src = ws.cell(r["source"], c, route.price.source)
            src.font = _font(BLUE, size=8)
            src.alignment = Alignment(wrap_text=True, vertical="top")
            if route.retailer_price is not None:
                self.flag_unverified(ws, _rel(c, r["retailer_price"]), f"{get_column_letter(c)}${r['status']}")
        if any(route.price.source for route in routes):
            ws.row_dimensions[r["source"]].height = 36
        self.status_ranges.append((ws.title, f"B{r['status']}:{last_col}{r['status']}"))
        self.flag_unverified(ws, f"B{r['price']}:{last_col}{r['price']}", f"B${r['status']}")
        self.rule_rows(ws, r["via"], r["source"], 1 + n)

        row += 1
        self.section(ws, row, f"Price waterfall per {unit}", 1 + n)
        row += 1
        put("mrp", "MRP (consumer price)")
        put("retailer_pays", "Retailer pays")
        put("retailer_margin", "Retailer margin")
        put("retailer_margin_pct", "Retailer margin, % of MRP")
        put("dist_margin", "Distributor margin")
        put("dist_margin_pct", "Distributor margin, % of its selling price")
        put("realization", "Brand realization per unit shipped (after scheme)", bold=True)
        put("net", "…net of output GST")
        put("landed", "Landed product cost")
        put("gross", "Gross contribution (on landed cost)", bold=True)
        put("gross_pct", "Gross margin")
        put("loaded", "Loaded offline cost (landed + allocations)")
        put("contribution", "Contribution after loaded cost", bold=True)
        put("contribution_pct", "Contribution margin")
        row += 1
        self.section(ws, row, "Checks", 1 + n)
        row += 1
        put("chk_contribution", "Contribution vs target")
        put("chk_gross", "Gross margin vs target")
        put("chk_band", "Retailer margin vs band")
        put("chk_dist", "Distributor margin")

        for i, route in enumerate(routes):
            c = 2 + i
            L = get_column_letter(c)

            def cell(key: str) -> str:
                return f"{L}{r[key]}"

            fx = {
                "mrp": ("=MRP", inr),
                "realization": (f"={cell('price')}*IF({cell('free')}>0,{cell('buy')}/({cell('buy')}+{cell('free')}),1)", inr),
                "net": (f"={cell('realization')}/(1+GSTRate)", inr),
                "retailer_pays": (f'=IF({cell("via")}="distributor",{cell("retailer_price")},{cell("realization")})', inr),
                "retailer_margin": (f"={cell('mrp')}-{cell('retailer_pays')}", inr),
                "retailer_margin_pct": (f"={cell('retailer_margin')}/{cell('mrp')}", PCT),
                "dist_margin": (f'=IF({cell("via")}="distributor",{cell("retailer_price")}-{cell("realization")},"–")', inr),
                "dist_margin_pct": (f'=IF({cell("via")}="distributor",{cell("dist_margin")}/{cell("retailer_price")},"–")', PCT),
                "landed": ("=LandedCost", inr),
                "gross": (f"={cell('net')}-{cell('landed')}", inr),
                "gross_pct": (f"={cell('gross')}/{cell('net')}", PCT),
                "loaded": ("=LoadedOffline", inr),
                "contribution": (f"={cell('net')}-{cell('loaded')}", inr),
                "contribution_pct": (f"={cell('contribution')}/{cell('net')}", PCT),
                "chk_contribution": (
                    f'=IF(COUNT(TgtOfflineContribution)=0,"–",IF({cell("contribution")}<TgtOfflineContribution-1E-9,"BELOW","OK"))',
                    None,
                ),
                "chk_gross": (f'=IF(COUNT(TgtGrossMargin)=0,"–",IF({cell("gross_pct")}<TgtGrossMargin-1E-9,"BELOW","OK"))', None),
                "chk_band": (
                    f'=IF(COUNT(TgtRetailerMarginMin,TgtRetailerMarginMax)<2,"–",'
                    f'IF(AND({cell("retailer_margin_pct")}>=TgtRetailerMarginMin-1E-9,'
                    f'{cell("retailer_margin_pct")}<=TgtRetailerMarginMax+1E-9),"IN BAND","OUTSIDE"))',
                    None,
                ),
                "chk_dist": (f'=IF({cell("via")}="distributor",IF({cell("dist_margin")}<0,"NEGATIVE","OK"),"–")', None),
            }
            for key, (text, fmt) in fx.items():
                cell_obj = self.formula(
                    ws, r[key], c, text, fmt, link=key in ("mrp", "landed", "loaded"), bold=key in ("realization", "gross", "contribution")
                )
                if key.startswith("chk_") or key.startswith("dist_margin") and route.via == "direct":
                    cell_obj.alignment = Alignment(horizontal="right")
            keys = {
                "realization": "realization",
                "net": "net_realization",
                "retailer_pays": "retailer_cost",
                "retailer_margin": "retailer_margin",
                "retailer_margin_pct": "retailer_margin_pct",
                "gross": "gross_contribution",
                "gross_pct": "gross_margin_pct",
                "loaded": "loaded_cost",
                "contribution": "contribution",
                "contribution_pct": "contribution_pct",
            }
            if route.via == "distributor":
                keys.update({"dist_margin": "distributor_margin", "dist_margin_pct": "distributor_margin_pct"})
            for key, metric in keys.items():
                self.track(f"routes.{route.id}.{metric}", ws, r[key], c)
        self.red_flags(ws, f"B{r['chk_contribution']}:{last_col}{r['chk_dist']}", "B" + str(r["chk_contribution"]), ("BELOW", "OUTSIDE", "NEGATIVE"))
        self.offline_rows = (ws.title, r, last_col)
        ws.freeze_panes = "B6"

    def online_sheet(self, channel: OnlineChannel) -> None:
        cfg, inr = self.cfg, self.inr
        ws = self.sheet(f"Online {channel.name}")
        unit = cfg.meta.unit
        self.title(ws, f"Online — {channel.name}: per-order economics", "Each column is one offer. Change prices, units and cost lines in blue.")
        n_offers = len(channel.offers)
        self.widths(ws, {"A": 56, **{get_column_letter(2 + i): 15 for i in range(max(n_offers, 6))}})

        row = 4
        self.section(ws, row, "Channel settings", 7)
        self.header(ws, row + 1, ["Setting", "Value", "Unit", "", "Status", "Source"])
        row += 2
        basis_row = row
        self.label(ws, row, 1, "Product cost basis (landed / loaded)")
        self.value(ws, row, 2, channel.unit_cost_basis)
        self.validation(ws, UNIT_COST_BASES, _rel(2, row))
        row += 1
        settings = [
            ("Irrecoverable GST on flagged cost lines", channel.irrecoverable_input_gst, PCT, "rate"),
            ("Shipping fee charged to the customer", channel.customer_shipping_fee, inr, "per order"),
            ("…charged on orders with item value below", channel.free_shipping_threshold, inr, "item value"),
        ]
        setting_rows = []
        first = row
        for label, param, fmt, unit_text in settings:
            self.label(ws, row, 1, label)
            self.value(ws, row, 2, param.value, fmt)
            self.label(ws, row, 3, unit_text, color=GREY)
            self.status_cells(ws, row, 5, param)
            setting_rows.append(row)
            row += 1
        self.status_ranges.append((ws.title, f"E{first}:E{row - 1}"))
        self.flag_unverified(ws, f"B{first}:B{row - 1}", f"$E{first}")
        uplift, fee, threshold = (_abs(2, r) for r in setting_rows)
        self.label(ws, row, 1, f"Product cost per {unit} used below")
        self.formula(ws, row, 2, f'=IF({_abs(2, basis_row)}="loaded",LoadedOnline,LandedCost)', inr)
        unit_cost = _abs(2, row)
        row += 2

        self.section(ws, row, "Per-order cost lines", 7)
        self.header(ws, row + 1, ["Cost line", "Kind", "Basis", "Amount", "Input GST applies (1/0)", "Status", "Source"])
        row += 2
        lines_first = row
        for line in channel.costs:
            self.label(ws, row, 1, line.name)
            self.value(ws, row, 2, line.kind)
            self.validation(ws, KINDS, _rel(2, row))
            self.value(ws, row, 3, line.basis)
            self.validation(ws, BASES, _rel(3, row))
            self.value(ws, row, 4, line.amount.value, PCT2 if line.basis == "pct_of_order_value" else inr)
            flag = self.value(ws, row, 5, 1 if line.input_gst else 0, "0")
            flag.alignment = Alignment(horizontal="center")
            self.validation(ws, ("0", "1"), _rel(5, row))
            self.status_cells(ws, row, 6, line.amount)
            row += 1
        lines_last = row - 1
        if channel.costs:
            self.rule_rows(ws, lines_first, lines_last, 7)
            self.status_ranges.append((ws.title, f"F{lines_first}:F{lines_last}"))
            self.flag_unverified(ws, f"D{lines_first}:D{lines_last}", f"$F{lines_first}")
        ws.cell(row, 1, "Basis: per_order = fixed per order · per_unit = × units · pct_of_order_value = share of what the customer pays").font = _font(GREY, italic=True)
        row += 2

        self.section(ws, row, "Offers — economics per order", 1 + max(n_offers, 1))
        row += 1
        self.header(ws, row, ["Per order"] + [o.name for o in channel.offers])
        row += 1
        cols = {o.id: 2 + i for i, o in enumerate(channel.offers)}
        r = {}

        def put(key: str, label: str, bold: bool = False) -> int:
            nonlocal row
            self.label(ws, row, 1, label, bold=bold)
            r[key] = row
            row += 1
            return r[key]

        put("units", "Units")
        put("price", "Price")
        put("status", "Price status")
        put("source", "Price source")
        for o in channel.offers:
            c = cols[o.id]
            self.value(ws, r["units"], c, o.units, "0")
            self.value(ws, r["price"], c, o.price.value, inr)
            st = ws.cell(r["status"], c, o.price.status)
            st.font = _font(BLUE)
            self.validation(ws, STATUSES, st.coordinate)
            src = ws.cell(r["source"], c, o.price.source)
            src.font = _font(BLUE, size=8)
            src.alignment = Alignment(wrap_text=True, vertical="top")
        if any(o.price.source for o in channel.offers):
            ws.row_dimensions[r["source"]].height = 36
        last_col = get_column_letter(1 + n_offers)
        self.status_ranges.append((ws.title, f"B{r['status']}:{last_col}{r['status']}"))
        test = ",".join(f'B${r["status"]}="{s}"' for s in UNVERIFIED)
        ws.conditional_formatting.add(f"B{r['price']}:{last_col}{r['price']}", FormulaRule(formula=[f"OR({test})"], fill=ASSUMPTION_FILL))

        put("mrp_value", "MRP value")
        put("discount", "Discount vs MRP")
        put("per_unit", f"Price per {unit}")
        put("fee", "Shipping fee paid by the customer")
        put("order_value", "Order value (what the customer pays)")
        put("net", "Net revenue (after output GST)")
        put("product", "Product cost")
        line_rows = [put(f"line.{line.id}", line.name) for line in channel.costs]
        put("total", "Total variable cost", bold=True)
        put("contribution", "Contribution per order", bold=True)
        put("per_unit_contribution", f"Contribution per {unit}")
        put("margin", "Contribution margin")
        put("marketing", "Acquisition cost in this order")
        put("cbm", "Break-even acquisition cost (contribution before marketing)")
        put("payback", "Orders needed to recover the acquisition cost")
        put("shipping", "Shipping cost in this order")
        put("be_ship", "Break-even shipping cost per order")
        put("check", "Check")

        kinds = f"$B${lines_first}:$B${lines_last}"
        for o in channel.offers:
            c = cols[o.id]
            L = get_column_letter(c)

            def cell(key: str) -> str:
                return f"{L}{r[key]}"

            fx = {
                "mrp_value": (f"={cell('units')}*MRP", inr),
                "discount": (f"=1-{cell('price')}/{cell('mrp_value')}", PCT),
                "per_unit": (f"={cell('price')}/{cell('units')}", inr),
                "fee": (f"=IF(AND({fee}>0,{cell('price')}<{threshold}),{fee},0)", inr),
                "order_value": (f"={cell('price')}+{cell('fee')}", inr),
                "net": (f"={cell('order_value')}/(1+GSTRate)", inr),
                "product": (f"={cell('units')}*{unit_cost}", inr),
            }
            for i, line in enumerate(channel.costs):
                lr = lines_first + i
                fx[f"line.{line.id}"] = (
                    f'=IF($C${lr}="per_order",$D${lr},IF($C${lr}="per_unit",$D${lr}*{cell("units")},'
                    f'$D${lr}*{cell("order_value")}))*IF($E${lr}=1,1+{uplift},1)',
                    inr,
                )
            line_sum = f"+SUM({L}{line_rows[0]}:{L}{line_rows[-1]})" if line_rows else ""
            fx["total"] = (f"={cell('product')}{line_sum}", inr)
            fx["contribution"] = (f"={cell('net')}-{cell('total')}", inr)
            fx["per_unit_contribution"] = (f"={cell('contribution')}/{cell('units')}", inr)
            fx["margin"] = (f"={cell('contribution')}/{cell('net')}", PCT)
            if line_rows:
                span = f"{L}{line_rows[0]}:{L}{line_rows[-1]}"
                fx["marketing"] = (f'=SUMIF({kinds},"marketing",{span})', inr)
                fx["shipping"] = (f'=SUMIF({kinds},"shipping",{span})', inr)
                fx["be_ship"] = (f'=IF(COUNTIF({kinds},"shipping")=0,"n/a",{cell("shipping")}+{cell("contribution")})', inr)
            else:
                fx["marketing"] = ("=0", inr)
                fx["shipping"] = ("=0", inr)
                fx["be_ship"] = ('="n/a"', None)
            fx["cbm"] = (f"={cell('contribution')}+{cell('marketing')}", inr)
            fx["payback"] = (f'=IF({cell("marketing")}=0,0,IF({cell("cbm")}>0,{cell("marketing")}/{cell("cbm")},"never"))', NUM)
            fx["check"] = (
                f'=IF({cell("cbm")}<=0,"LOSS before marketing",IF(AND(COUNT(TgtOnlineContribution)=1,'
                f'{cell("contribution")}<TgtOnlineContribution-1E-9),"BELOW target","OK"))',
                None,
            )
            for key, (text, fmt) in fx.items():
                bold = key in ("total", "contribution")
                self.formula(ws, r[key], c, text, fmt, bold=bold)
            base = f"online.{channel.id}.{o.id}"
            for key, metric in (
                ("mrp_value", "mrp_value"),
                ("discount", "discount_vs_mrp"),
                ("per_unit", "price_per_unit"),
                ("fee", "shipping_fee_income"),
                ("order_value", "order_value"),
                ("net", "net_revenue"),
                ("product", "product_cost"),
                ("total", "total_cost"),
                ("contribution", "contribution"),
                ("per_unit_contribution", "contribution_per_unit"),
                ("margin", "contribution_margin"),
                ("marketing", "marketing"),
                ("cbm", "contribution_before_marketing"),
                ("payback", "cac_payback_orders"),
                ("be_ship", "break_even_shipping"),
            ):
                self.track(f"{base}.{metric}", ws, r[key], c)
            for line in channel.costs:
                self.track(f"{base}.line.{line.id}", ws, r[f"line.{line.id}"], c)
        self.red_flags(ws, f"B{r['check']}:{last_col}{r['check']}", f"B{r['check']}", ("LOSS before marketing", "BELOW target"))
        for c in cols.values():
            ws.cell(r["check"], c).alignment = Alignment(horizontal="right")
        self.online_layout = getattr(self, "online_layout", {})
        self.online_layout[channel.id] = {
            "sheet": ws.title,
            "units_row": r["units"],
            "price_row": r["price"],
            "check_row": r["check"],
            "cols": cols,
            "fee": fee,
            "threshold": threshold,
            "n_offers": n_offers,
        }
        ws.freeze_panes = "B4"
        self.ladder_sheet(channel)

    def ladder_sheet(self, channel: OnlineChannel) -> None:
        inr = self.inr
        layout = self.online_layout[channel.id]
        src = layout["sheet"]
        ws = self.sheet(f"Ladder {channel.name}")
        self.title(
            ws,
            f"Price ladder — {channel.name}",
            "Is every bigger pack cheaper than buying the same quantity as smaller packs? Same-size multiples are tested here.",
        )
        by_id = {o.id: o for o in channel.offers}
        self.ladder_rows = getattr(self, "ladder_rows", [])
        row = 4
        widths_needed = 0
        for card in channel.price_cards:
            offers = sorted((by_id[i] for i in card.offer_ids), key=lambda o: o.units)
            n = len(offers)
            width = 4 + n + 5
            widths_needed = max(widths_needed, width)
            self.section(ws, row, f"Price card: {card.name}", width)
            row += 1
            labels = ["Offer", "Units", "Price", f"Price per {self.cfg.meta.unit}"]
            labels += [f"Buy as {o.name}" for o in offers]
            labels += ["Cheapest same-size alternative", "Customer pays for the pack", "Customer pays for the alternative", "Pack saves the customer", "Status"]
            self.header(ws, row, labels)
            ws.row_dimensions[row].height = 42
            row += 1
            first = row
            for i, offer in enumerate(offers):
                rr = first + i
                oc = layout["cols"][offer.id]
                self.label(ws, rr, 1, offer.name)
                self.formula(ws, rr, 2, f"={_ext(src, oc, layout['units_row'])}", "0", link=True)
                self.formula(ws, rr, 3, f"={_ext(src, oc, layout['price_row'])}", inr, link=True)
                self.formula(ws, rr, 4, f"=C{rr}/B{rr}", inr)
                for j in range(n):
                    rj = first + j
                    self.formula(
                        ws,
                        rr,
                        5 + j,
                        f'=IF(AND($B${rj}<$B{rr},MOD($B{rr},$B${rj})=0),$C${rj}*$B{rr}/$B${rj},"")',
                        inr,
                    )
                span = f"{_rel(5, rr)}:{_rel(4 + n, rr)}"
                c_alt, c_pack, c_alt_total, c_save, c_status = (5 + n + k for k in range(5))
                alt = _rel(c_alt, rr)
                self.formula(ws, rr, c_alt, f'=IF(COUNT({span})=0,"–",MIN({span}))', inr)
                fee, thr = layout["fee"], layout["threshold"]
                fee_src = f"'{src}'!{fee}"
                thr_src = f"'{src}'!{thr}"
                self.formula(ws, rr, c_pack, f"=C{rr}+IF(AND({fee_src}>0,C{rr}<{thr_src}),{fee_src},0)", inr)
                self.formula(
                    ws,
                    rr,
                    c_alt_total,
                    f'=IF(ISNUMBER({alt}),{alt}+IF(AND({fee_src}>0,{alt}<{thr_src}),{fee_src},0),"–")',
                    inr,
                )
                self.formula(
                    ws,
                    rr,
                    c_save,
                    f'=IF(ISNUMBER({_rel(c_alt_total, rr)}),{_rel(c_alt_total, rr)}-{_rel(c_pack, rr)},"–")',
                    inr,
                )
                save = _rel(c_save, rr)
                self.formula(
                    ws,
                    rr,
                    c_status,
                    f'=IF(ISNUMBER({save}),IF({save}<=1E-9,"DOMINATED","OK"),"smallest")',
                    bold=True,
                )
                key = f"ladder.{channel.id}.{card.id}.{offer.id}"
                self.track(f"{key}.customer_total", ws, rr, c_pack)
                self.track(f"{key}.alt_pure", ws, rr, c_alt)
                self.track(f"{key}.dominated", ws, rr, c_status)
                self.ladder_rows.append((key, ws.title, _rel(c_status, rr)))
            status_col = get_column_letter(5 + n + 4)
            self.red_flags(ws, f"{status_col}{first}:{status_col}{first + n - 1}", f"{status_col}{first}", ("DOMINATED",))
            row = first + n + 1
        self.widths(ws, {"A": 24, "B": 8, "C": 13, "D": 13, **{get_column_letter(c): 15 for c in range(5, widths_needed + 1)}})
        ws.cell(row, 1, "Pairs that cannot make the same quantity stay blank. Run the engine for mixed combinations such as 6 + 3 + 3.").font = _font(GREY, italic=True)

    def scenario_sheet(self) -> None:
        cfg, inr = self.cfg, self.inr
        big = self.inr_big
        ws = self.sheet("Scenarios")
        self.title(ws, "Period scenarios — P&L for a month (or any period)", "If channel-mix shares are filled in they set the realization; otherwise the blended input is used.")
        self.widths(ws, {"A": 46, "B": 16, "C": 16, "D": 13, "E": 13, "F": 50})
        self.scenario_checks = []
        row = 4
        for s in cfg.scenarios:
            row = self._scenario_block(ws, s, row, inr, big) + 2
        ws.freeze_panes = "A4"

    def _scenario_block(self, ws, s: Scenario, row: int, inr: str, big: str) -> int:
        unit = self.cfg.meta.unit
        self.section(ws, row, f"Scenario: {s.name}", 6)
        self.header(ws, row + 1, ["Input", "Value", "", "", "Status", "Source"])
        row += 2
        units_row = row
        self.label(ws, row, 1, f"Units ({unit}s) in the period")
        self.value(ws, row, 2, s.units.value, INT)
        self.status_cells(ws, row, 5, s.units)
        row += 1
        blended_row = row
        self.label(ws, row, 1, f"Blended realization per {unit} (used when no mix is entered)")
        bp = s.blended_realization
        self.value(ws, row, 2, bp.value if bp else None, inr)
        if bp:
            self.status_cells(ws, row, 5, bp)
        else:
            self.validation(ws, STATUSES, _rel(5, row))
        self.status_ranges.append((ws.title, f"E{units_row}:E{row}"))
        self.flag_unverified(ws, f"B{units_row}:B{row}", f"$E{units_row}")
        row += 2

        self.header(ws, row, ["Channel mix", "Share", f"Realization per {unit}", "", "Status", "Source"])
        row += 1
        mix_first = row
        for m in s.mix:
            self.label(ws, row, 1, m.name)
            self.value(ws, row, 2, m.share.value, PCT)
            self.value(ws, row, 3, m.realization.value, inr)
            self.status_cells(ws, row, 5, m.share)
            row += 1
        for _ in range(max(EXTRA_ROWS, 4 - len(s.mix))):
            self.value(ws, row, 1, None)
            self.value(ws, row, 2, None, PCT)
            self.value(ws, row, 3, None, inr)
            self.validation(ws, STATUSES, _rel(5, row))
            row += 1
        mix_last = row - 1
        self.rule_rows(ws, mix_first, mix_last, 6)
        self.status_ranges.append((ws.title, f"E{mix_first}:E{mix_last}"))
        self.flag_unverified(ws, f"B{mix_first}:C{mix_last}", f"$E{mix_first}")
        shares = f"$B${mix_first}:$B${mix_last}"
        reals = f"$C${mix_first}:$C${mix_last}"
        self.label(ws, row, 1, "Mix shares add up to")
        self.formula(ws, row, 2, f"=SUM({shares})", PCT)
        share_sum = _abs(2, row)
        self.formula(ws, row, 3, f'=IF({share_sum}=0,"no mix entered",IF(ABS({share_sum}-1)<1E-9,"OK","CHECK: not 100%"))')
        row += 2

        self.header(ws, row, ["Operating cost line", "Amount for the period", "Kind", "", "Status", "Source"])
        row += 1
        opex_first = row
        for line in s.opex:
            self.label(ws, row, 1, line.name)
            self.value(ws, row, 2, line.amount.value, big)
            self.value(ws, row, 3, line.kind)
            self.validation(ws, KINDS, _rel(3, row))
            self.status_cells(ws, row, 5, line.amount)
            row += 1
        for _ in range(EXTRA_ROWS):
            self.value(ws, row, 1, None)
            self.value(ws, row, 2, None, big)
            self.value(ws, row, 3, None)
            self.validation(ws, KINDS, _rel(3, row))
            self.validation(ws, STATUSES, _rel(5, row))
            row += 1
        opex_last = row - 1
        self.rule_rows(ws, opex_first, opex_last, 6)
        ws.cell(row, 1, "Blank rows take extra cost lines; mark marketing spend with kind = marketing.").font = _font(GREY, italic=True)
        row += 1
        self.status_ranges.append((ws.title, f"E{opex_first}:E{opex_last}"))
        self.flag_unverified(ws, f"B{opex_first}:B{opex_last}", f"$E{opex_first}")
        amounts = f"$B${opex_first}:$B${opex_last}"
        kinds = f"$C${opex_first}:$C${opex_last}"
        row += 1

        self.header(ws, row, ["Result", "Amount", "", "", "", ""])
        row += 1
        u, b = _abs(2, units_row), _abs(2, blended_row)
        rows = {}
        results = [
            ("realization", f"Realization per {unit} used", f"=IF({share_sum}>0,SUMPRODUCT({shares},{reals})/{share_sum},{b})", inr),
            ("revenue", "Revenue (net of output GST)", f"={u}*B{{realization}}/(1+GSTRate)", big),
            ("landed_total", "Landed product cost", f"={u}*LandedCost", big),
            ("gross_contribution", "Gross contribution", "=B{revenue}-B{landed_total}", big),
            ("gross_margin_pct", "Gross margin", '=IF(B{revenue}=0,0,B{gross_contribution}/B{revenue})', PCT),
            ("opex_total", "Total operating costs", f"=SUM({amounts})", big),
            ("operating_contribution", "Operating contribution", "=B{gross_contribution}-B{opex_total}", big),
            ("operating_margin_pct", "Operating margin", '=IF(B{revenue}=0,0,B{operating_contribution}/B{revenue})', PCT),
            ("opex_per_unit", f"Operating costs per {unit}", f"=B{{opex_total}}/{u}", inr),
            ("implied_cost_per_unit", f"Implied all-in cost per {unit}", "=LandedCost+B{opex_per_unit}", inr),
            (
                "implied_cost_excl_marketing",
                f"Implied cost per {unit} excluding marketing",
                f'=LandedCost+(B{{opex_total}}-SUMIF({kinds},"marketing",{amounts}))/{u}',
                inr,
            ),
        ]
        for key, label, text, fmt in results:
            rows[key] = row
            self.label(ws, row, 1, label, bold=key == "operating_contribution")
            self.formula(ws, row, 2, text.format(**rows), fmt, bold=key == "operating_contribution")
            self.track(f"scenarios.{s.id}.{key}", ws, row, 2)
            row += 1
        self.track(f"scenarios.{s.id}.units", ws, units_row, 2)
        self.scenario_checks.append(
            (ws.title, _abs(2, rows["operating_contribution"]), _abs(2, rows["implied_cost_per_unit"]))
        )
        return row

    def checks_sheet(self) -> None:
        ws = self.sheet("Checks")
        self.title(ws, "Checks — what needs attention", "Counts update as you change inputs. Anything other than OK deserves a decision.")
        self.widths(ws, {"A": 62, "B": 10, "C": 14, "D": 40})
        self.header(ws, 4, ["Check", "Count", "Status", "Where to look"])
        row = 5
        checks: list[tuple[str, str, str, str]] = []
        checks.append(("PLANNING_COST_OUT_OF_RANGE", "All-in cost outside the stated planning range", '=IF(PlanCheck="NO",1,0)', "Inputs"))
        if self.cfg.routes:
            sheet, rows, last_col = self.offline_rows
            q = f"'{sheet}'!"

            def span(key: str) -> str:
                return f"{q}B{rows[key]}:{last_col}{rows[key]}"

            checks += [
                ("NEGATIVE_DISTRIBUTOR_MARGIN", "Routes with a negative distributor margin", f'=COUNTIF({span("chk_dist")},"NEGATIVE")', sheet),
                ("OFFLINE_BELOW_TARGET", "Routes below the offline contribution target", f'=COUNTIF({span("chk_contribution")},"BELOW")', sheet),
                ("GROSS_MARGIN_BELOW_TARGET", "Routes below the gross-margin target", f'=COUNTIF({span("chk_gross")},"BELOW")', sheet),
                ("RETAILER_MARGIN_OUTSIDE_BAND", "Routes with retailer margin outside the band", f'=COUNTIF({span("chk_band")},"OUTSIDE")', sheet),
            ]
        for channel in self.cfg.online:
            lay = self.online_layout[channel.id]
            q = f"'{lay['sheet']}'!"
            span = f"{q}B{lay['check_row']}:{get_column_letter(1 + lay['n_offers'])}{lay['check_row']}"
            checks += [
                (f"ONLINE_LOSS_BEFORE_MARKETING:{channel.id}", f"{channel.name}: offers losing money before marketing", f'=COUNTIF({span},"LOSS before marketing")', lay["sheet"]),
                (f"ONLINE_BELOW_TARGET:{channel.id}", f"{channel.name}: offers below the contribution target", f'=COUNTIF({span},"BELOW target")', lay["sheet"]),
            ]
            dominated = [ref for key, sheet, ref in self.ladder_rows if key.startswith(f"ladder.{channel.id}.")]
            if dominated:
                ladder_sheet = [sheet for key, sheet, ref in self.ladder_rows if key.startswith(f"ladder.{channel.id}.")][0]
                terms = "+".join(f"IF('{ladder_sheet}'!{ref}=\"DOMINATED\",1,0)" for ref in dominated)
                checks.append((f"LADDER_DOMINATED:{channel.id}", f"{channel.name}: packs that are no cheaper than smaller packs", f"={terms}", ladder_sheet))
        if self.cfg.scenarios:
            loss = "+".join(f"IF('{sheet}'!{oc}<0,1,0)" for sheet, oc, _ in self.scenario_checks)
            above = "+".join(f"IF(COUNT(PlanMax)=0,0,IF('{sheet}'!{ic}>PlanMax+1E-9,1,0))" for sheet, _, ic in self.scenario_checks)
            checks += [
                ("SCENARIO_OPERATING_LOSS", "Scenarios with an operating loss", f"={loss}", "Scenarios"),
                ("SCENARIO_COST_ABOVE_PLANNING", "Scenarios whose implied cost exceeds the planning range", f"={above}", "Scenarios"),
            ]
        for code, label, text, where in checks:
            self.label(ws, row, 1, label)
            self.formula(ws, row, 2, text, "0").alignment = Alignment(horizontal="center")
            self.formula(ws, row, 3, f'=IF(B{row}=0,"OK","REVIEW")', bold=True).alignment = Alignment(horizontal="center")
            self.label(ws, row, 4, where, color=GREY)
            self.check_cells[code] = (ws.title, _rel(2, row))
            row += 1
        ws.conditional_formatting.add(f"C5:C{row - 1}", FormulaRule(formula=['C5="REVIEW"'], font=Font(name=FONT, bold=True, color=RED)))
        ws.freeze_panes = "A5"

    def readme_sheet(self, ws) -> None:
        cfg = self.cfg
        self.widths(ws, {"A": 30, "B": 16, "C": 90})
        self.title(ws, f"{cfg.meta.name} — unit economics model", "Generated by factory/unit_economics from a TOML config.")
        row = 3
        if cfg.meta.confidential:
            ws.cell(row, 1, "CONFIDENTIAL — contains cost and margin data. Keep this file private.").font = _font(RED, bold=True, size=11)
            row += 1
        if cfg.meta.notes:
            ws.cell(row, 1, cfg.meta.notes).font = _font(GREY, italic=True)
            row += 1
        row += 1
        self.section(ws, row, "How to use", 3)
        row += 1
        steps = [
            "Change only blue cells. Every other number is a formula and recalculates immediately.",
            "A yellow fill means the input is still an assumption (status report, assumption or unspecified). "
            "When a quote, invoice or actual arrives, type it in, set Status to quote or actual and note the Source; the fill clears.",
            "Inputs: price, landed cost, GST, cost build-up and targets. Offline: routes and partner margins. "
            "Online: per-order economics for each offer. Ladder: is each bigger pack cheaper than smaller ones? "
            "Scenarios: a period P&L. Checks: everything that needs a decision.",
            "To add or remove an offer, route, cost line or scenario, edit the TOML config and regenerate: "
            "python -m unit_economics workbook CONFIG --out FILE.xlsx",
            "Ladder sheets test same-size multiples (e.g. 2 × 6-pack). The engine also tests mixed combinations; "
            "the two agree whenever each pack size divides the larger ones.",
        ]
        for i, text in enumerate(steps, 1):
            cell = ws.cell(row, 1, f"{i}. {text}")
            cell.font = _font()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            ws.row_dimensions[row].height = 14 * (len(text) // 105 + 1) + 4
            row += 1
        row += 1
        self.section(ws, row, "Legend", 3)
        row += 1
        legend = [
            ("123", _font(BLUE), None, "Input you can change"),
            ("123", _font(BLUE), ASSUMPTION_FILL, "Input that is still an assumption"),
            ("123", _font(BLACK), None, "Formula"),
            ("123", _font(GREEN), None, "Link to another sheet"),
        ]
        for sample, font, fill, text in legend:
            cell = ws.cell(row, 1, sample)
            cell.font = font
            if fill:
                cell.fill = fill
            ws.cell(row, 2, text).font = _font()
            row += 1
        row += 1
        self.section(ws, row, "Evidence status of inputs", 3)
        row += 1
        self.header(ws, row, ["Status", "Inputs", "Meaning"])
        row += 1
        meanings = {
            "report": "Taken from a planning report; not yet checked against a quote or actual",
            "assumption": "Analyst assumption; replace when data exists",
            "research": "From a public source (see Source); re-check before relying on it",
            "quote": "Supplier or partner quotation",
            "actual": "From invoices or real operating data",
            "unspecified": "No status recorded",
        }
        first = row
        for status in STATUSES:
            ws.cell(row, 1, status).font = _font()
            terms = "+".join(f"COUNTIF('{sheet}'!{rng},\"{status}\")" for sheet, rng in self.status_ranges)
            self.formula(ws, row, 2, f"={terms}", "0")
            ws.cell(row, 3, meanings[status]).font = _font(GREY)
            row += 1
        ws.cell(row, 1, "Total").font = _font(bold=True)
        self.formula(ws, row, 2, f"=SUM(B{first}:B{row - 1})", "0", bold=True)
        self.status_count_cells = {status: (ws.title, _rel(2, first + i)) for i, status in enumerate(STATUSES)}


def build_workbook(config: Config) -> tuple[Workbook, dict[str, tuple[str, str]], dict[str, tuple[str, str]], dict[str, tuple[str, str]]]:
    """Build the workbook.

    Returns (workbook, result cells, check-count cells, status-count cells);
    each map gives (sheet title, cell coordinate) per key.
    """
    builder = _Builder(config)
    wb = builder.build()
    return wb, builder.cells, builder.check_cells, builder.status_count_cells


def save_workbook(config: Config, path: str | Path) -> dict[str, tuple[str, str]]:
    wb, cells, _, _ = build_workbook(config)
    wb.save(str(path))
    return cells
